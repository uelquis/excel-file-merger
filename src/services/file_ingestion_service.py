"""Service that ingests excel files and merges worksheet rows by sheet name."""

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet

from domain.cell_content import CellContent
from domain.merged_workbook import MergedWorkbook, MergedWorksheet
from events.event_bus import EventBus, FilesLoaded, RowsMerged, ValidationWarning

# Per-sheet read result: title, header cells, non-blank data rows, and how
# many blank rows were dropped while reading that sheet.
_ReadSheetResult = tuple[str, list[CellContent], list[list[CellContent]], int]


class FileIngestionService:
    """Loads user-selected files/folders and merges their worksheets by name.

    Merging is order-independent: every worksheet with the same name across
    all ingested files is merged into a single result worksheet (for example
    file1: a,b + file2: b,a -> a1+a2, b1+b2). Validation is soft: any
    inconsistency publishes a ValidationWarning event and the merge proceeds.

    Cell background colors (RGB solid fills) are captured into CellContent so
    they can be preserved on output. Blank rows are dropped during ingestion
    and blank columns are dropped after merging, so the merged result never
    contains fully-empty rows or columns. "Blank" always means no value; a
    background color never counts as content.

    Composed with an EventBus instead of owning observers directly.
    """

    EXCEL_FILE_SUFFIX = ".xlsx"

    def __init__(self, event_bus: EventBus) -> None:
        """Create the ingestion service.

        Args:
            event_bus: Bus used to publish FilesLoaded, RowsMerged and
                ValidationWarning events.
        """
        self._event_bus = event_bus

    def collect_excel_file_paths(self, selected_paths: list[Path]) -> list[Path]:
        """Expand the user selection into an ordered list of .xlsx files.

        Folders are expanded into the .xlsx files they contain; individual
        files are kept as-is. Selection order is preserved and duplicates
        are removed.

        Args:
            selected_paths: Files and/or folders chosen by the user.

        Returns:
            Ordered, deduplicated list of .xlsx file paths.
        """
        collected_file_paths: list[Path] = []
        for selected_path in selected_paths:
            if selected_path.is_dir():
                collected_file_paths.extend(
                    sorted(selected_path.glob(f"*{self.EXCEL_FILE_SUFFIX}"))
                )
            elif selected_path.suffix.lower() == self.EXCEL_FILE_SUFFIX:
                collected_file_paths.append(selected_path)
            else:
                self._event_bus.publish(
                    ValidationWarning(
                        f"Skipped '{selected_path.name}': not a {self.EXCEL_FILE_SUFFIX} file"
                    )
                )

        unique_file_paths: list[Path] = []
        for file_path in collected_file_paths:
            if file_path not in unique_file_paths:
                unique_file_paths.append(file_path)

        self._event_bus.publish(FilesLoaded(file_paths=unique_file_paths))
        return unique_file_paths

    def merge_worksheet_rows(self, file_paths: list[Path]) -> MergedWorkbook:
        """Merge the rows of all worksheets, grouped by sheet name.

        For each unique sheet name across all files, a MergedWorksheet is
        created holding the header from the first file containing that sheet
        and the data rows appended from every file. Result sheet order
        follows the first file; names appearing only in later files are
        appended in order of appearance.

        Blank rows are dropped while reading each sheet; blank columns are
        dropped per merged worksheet once all files are merged (before schema
        generation, so removed columns never reach the editor).

        Soft validations published as ValidationWarning events:
        - a sheet name present in only some files (still merged);
        - same sheet name with different header values (first file wins);
        - a file that could not be opened (skipped);
        - counts of dropped blank rows/columns per sheet.

        Args:
            file_paths: Ordered list of .xlsx files to merge.

        Returns:
            MergedWorkbook with one merged worksheet per unique sheet name
            (schemas are not generated yet).

        Raises:
            ValueError: If no excel file was provided.
        """
        if not file_paths:
            raise ValueError("No excel files provided to merge")

        merged_workbook = MergedWorkbook()
        sheet_presence_by_title: dict[str, list[str]] = {}
        blank_rows_dropped_by_title: dict[str, int] = {}

        for file_path in file_paths:
            sheets_in_file = self._read_file_sheets(file_path)
            if sheets_in_file is None:
                continue

            for sheet_title, header_row, data_rows, blank_rows_dropped in sheets_in_file:
                sheet_presence_by_title.setdefault(sheet_title, []).append(file_path.name)
                blank_rows_dropped_by_title[sheet_title] = (
                    blank_rows_dropped_by_title.get(sheet_title, 0) + blank_rows_dropped
                )
                existing_merged_worksheet = merged_workbook.get_worksheet_by_title(sheet_title)

                if existing_merged_worksheet is None:
                    merged_workbook.add_worksheet(
                        MergedWorksheet(
                            sheet_title=sheet_title,
                            header_row=header_row,
                            data_rows=data_rows,
                        )
                    )
                    continue

                if self._header_values(header_row) != self._header_values(
                    existing_merged_worksheet.header_row
                ):
                    self._event_bus.publish(
                        ValidationWarning(
                            f"Sheet '{sheet_title}' in '{file_path.name}' has a different header "
                            "than the first file containing it; keeping the first header",
                            sheet_title=sheet_title,
                        )
                    )
                existing_merged_worksheet.append_data_rows(data_rows)

        blank_columns_dropped_by_title = {
            worksheet.sheet_title: self._drop_blank_columns(worksheet)
            for worksheet in merged_workbook.worksheets
        }

        self._warn_about_partially_present_sheets(sheet_presence_by_title, file_paths)
        self._warn_about_dropped_blanks(
            blank_rows_dropped_by_title, blank_columns_dropped_by_title
        )

        self._event_bus.publish(RowsMerged(merged_workbook=merged_workbook))
        return merged_workbook

    def _read_file_sheets(self, file_path: Path) -> list[_ReadSheetResult] | None:
        """Open one workbook, read all sheets' cells, then close it.

        Loaded in normal (non read-only) mode so cell background fills are
        reliably available for every cell.

        Args:
            file_path: Path of the .xlsx file to read.

        Returns:
            List of (sheet_title, header_row, data_rows, blank_rows_dropped)
            per worksheet, or None if the file could not be opened.
        """
        try:
            workbook = load_workbook(file_path)
        except Exception as load_error:
            self._event_bus.publish(
                ValidationWarning(
                    f"Skipped '{file_path.name}': could not be opened ({load_error})"
                )
            )
            return None

        try:
            return [
                self._read_header_and_data_rows(source_worksheet)
                for source_worksheet in workbook.worksheets
            ]
        finally:
            workbook.close()

    def _read_header_and_data_rows(
        self, source_worksheet: Worksheet
    ) -> _ReadSheetResult:
        """Split a source worksheet into title, header, non-blank data rows.

        The first row is always treated as the header. Every subsequent row
        that has no value in any cell is dropped as blank and counted.

        Args:
            source_worksheet: Worksheet to read from top to bottom.

        Returns:
            Tuple of (sheet_title, header_row, data_rows, blank_rows_dropped);
            header_row is empty when the worksheet has no rows at all.
        """
        rows_as_cells: list[list[CellContent]] = [
            [self._cell_to_content(cell) for cell in row]
            for row in source_worksheet.iter_rows()
        ]
        if not rows_as_cells:
            return source_worksheet.title, [], [], 0

        header_row = rows_as_cells[0]
        data_rows: list[list[CellContent]] = []
        blank_rows_dropped = 0
        for cells in rows_as_cells[1:]:
            if any(cell.has_value() for cell in cells):
                data_rows.append(cells)
            else:
                blank_rows_dropped += 1

        return source_worksheet.title, header_row, data_rows, blank_rows_dropped

    def _cell_to_content(self, cell: Cell) -> CellContent:
        """Capture one cell's value and RGB solid background color.

        Args:
            cell: openpyxl cell to read.

        Returns:
            CellContent holding the value and the captured background color.
        """
        return CellContent(
            value=cell.value, background_color=self._extract_background_color(cell)
        )

    def _extract_background_color(self, cell: Cell) -> str | None:
        """Extract the RGB hex of a solid fill, ignoring theme/indexed colors.

        Args:
            cell: openpyxl cell whose fill is inspected.

        Returns:
            The RGB hex string (for example 'FFFF0000') when the cell has a
            solid fill with an explicit RGB color; None otherwise (no fill,
            non-solid fill, or theme/indexed color).
        """
        fill = cell.fill
        if fill is None or fill.fill_type != "solid":
            return None
        start_color = fill.start_color
        if start_color is None:
            return None
        if start_color.type == "rgb" and isinstance(start_color.rgb, str):
            return start_color.rgb
        return None

    def _drop_blank_columns(self, merged_worksheet: MergedWorksheet) -> int:
        """Remove columns that are empty in both header and every data row.

        A column is blank only when its header cell has no value AND all of
        its data cells have no value; any column with content (header or data)
        is kept. Background color never keeps a column alive. Rows are
        rewritten so remaining column indices stay consistent.

        Args:
            merged_worksheet: Worksheet to trim in place.

        Returns:
            Number of blank columns removed.
        """
        header_row = merged_worksheet.header_row
        data_rows = merged_worksheet.data_rows
        column_count = self._widest_row_length(header_row, data_rows)

        kept_column_indices: list[int] = []
        for column_index in range(column_count):
            if self._column_has_value(column_index, header_row, data_rows):
                kept_column_indices.append(column_index)

        dropped_column_count = column_count - len(kept_column_indices)
        if dropped_column_count == 0:
            return 0

        merged_worksheet.header_row = [
            header_row[column_index]
            for column_index in kept_column_indices
            if column_index < len(header_row)
        ]
        merged_worksheet.data_rows = [
            [
                data_row[column_index]
                for column_index in kept_column_indices
                if column_index < len(data_row)
            ]
            for data_row in data_rows
        ]
        return dropped_column_count

    def _widest_row_length(
        self, header_row: list[CellContent], data_rows: list[list[CellContent]]
    ) -> int:
        """Return the number of columns spanned by the header and data rows.

        Args:
            header_row: Header cells.
            data_rows: Data rows (possibly ragged).

        Returns:
            Width of the widest row (0 when everything is empty).
        """
        widest_data_row_length = max((len(row) for row in data_rows), default=0)
        return max(len(header_row), widest_data_row_length)

    def _column_has_value(
        self,
        column_index: int,
        header_row: list[CellContent],
        data_rows: list[list[CellContent]],
    ) -> bool:
        """Return True when any cell in the given column carries a value.

        Args:
            column_index: Zero-based column to inspect.
            header_row: Header cells.
            data_rows: Data rows (possibly ragged).
        """
        if column_index < len(header_row) and header_row[column_index].has_value():
            return True
        return any(
            column_index < len(data_row) and data_row[column_index].has_value()
            for data_row in data_rows
        )

    def _header_values(self, header_row: list[CellContent]) -> list:
        """Extract just the values of a header row for mismatch comparison.

        Colors may differ across files without that being a header mismatch,
        so only values are compared.

        Args:
            header_row: Header cells.

        Returns:
            List of raw header values.
        """
        return [cell.value for cell in header_row]

    def _warn_about_partially_present_sheets(
        self, sheet_presence_by_title: dict[str, list[str]], file_paths: list[Path]
    ) -> None:
        """Warn about sheet names that are not present in every ingested file.

        Args:
            sheet_presence_by_title: Sheet title -> names of files containing it.
            file_paths: All files that were merged.
        """
        total_file_count = len(file_paths)
        for sheet_title, present_file_names in sheet_presence_by_title.items():
            if len(present_file_names) < total_file_count:
                self._event_bus.publish(
                    ValidationWarning(
                        f"Sheet '{sheet_title}' only found in {', '.join(present_file_names)}; "
                        "merged with the rows available",
                        sheet_title=sheet_title,
                    )
                )

    def _warn_about_dropped_blanks(
        self,
        blank_rows_dropped_by_title: dict[str, int],
        blank_columns_dropped_by_title: dict[str, int],
    ) -> None:
        """Report how many blank rows/columns were removed per sheet.

        Args:
            blank_rows_dropped_by_title: Sheet title -> blank rows dropped.
            blank_columns_dropped_by_title: Sheet title -> blank columns dropped.
        """
        for sheet_title, dropped_row_count in blank_rows_dropped_by_title.items():
            if dropped_row_count > 0:
                self._event_bus.publish(
                    ValidationWarning(
                        f"Sheet '{sheet_title}': removed {dropped_row_count} blank row(s)",
                        sheet_title=sheet_title,
                    )
                )
        for sheet_title, dropped_column_count in blank_columns_dropped_by_title.items():
            if dropped_column_count > 0:
                self._event_bus.publish(
                    ValidationWarning(
                        f"Sheet '{sheet_title}': removed {dropped_column_count} blank column(s)",
                        sheet_title=sheet_title,
                    )
                )
